"""
export_pptx.py - build output/Readings5_Presentation.pptx (16:9) with the same
slides, titles and key points as the Streamlit app. Every demo is run at its
default slider values and saved as a PNG; speaker notes go in the notes pane.

Run:  python export_pptx.py
"""
import json
import time

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import config
import content
from demos import actuator_net, domain_rand, kine2go, pipeline, sdk_demo
from demos.style import plt, save_figure

SLIDE_W, SLIDE_H = 13.333, 7.5          # inches (16:9)
MARGIN = 0.5
BAR_H = 1.15                            # title bar height
TITLE_PT = 32                           # minimum title size required: 32 pt
BODY_PT = 18                            # body text (never below 16 pt)
MIN_BODY_PT = 16
FONT = "Calibri"


def rgb(hex_color):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ----------------------------------------------------------------------------
# Text fitting: a simple estimate so text never overflows its box
# ----------------------------------------------------------------------------
def estimate_lines(text, width_in, pt, indent_in=0.0):
    """Rough line count for `text` at `pt` in a box `width_in` wide
    (Calibri averages about 0.5 em per character)."""
    chars_per_line = max(1, int((width_in - indent_in - 0.2) * 72 / (pt * 0.5)))
    words, lines, cur = text.split(), 1, 0
    for w in words:
        if cur and cur + 1 + len(w) > chars_per_line:
            lines, cur = lines + 1, len(w)
        else:
            cur += (1 if cur else 0) + len(w)
    return lines


def text_height_in(paragraphs, width_in, pt, space_after_pt, indent_in=0.0):
    lines = sum(estimate_lines(p, width_in, pt, indent_in) for p in paragraphs)
    return (lines * pt * 1.2 + len(paragraphs) * space_after_pt) / 72 + 0.15


def fit_size(paragraphs, width_in, height_in, start_pt=BODY_PT, space_after_pt=8, indent_in=0.3):
    """Largest font size (>= 16 pt) at which the paragraphs fit the box."""
    pt = start_pt
    while pt > MIN_BODY_PT and text_height_in(paragraphs, width_in, pt, space_after_pt, indent_in) > height_in:
        pt -= 1
    fits = text_height_in(paragraphs, width_in, pt, space_after_pt, indent_in) <= height_in
    return pt, fits


FIT_REPORT = []   # (slide title, box name, font size, fits?)


# ----------------------------------------------------------------------------
# Shape helpers
# ----------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _set_bullet(paragraph, char="•", indent_in=0.3, color=config.ACCENT):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent_in))))
    pPr.set("indent", str(-Emu(Inches(indent_in))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": color.lstrip("#")})
    buClr.append(srgb)
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def add_text(slide, x, y, w, h, paragraphs, pt=BODY_PT, color=config.TEXT_DARK, bold=False,
             bullet=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=8, name="text",
             title_for_report="", autofit=True, numbered=False):
    if autofit:
        pt, fits = fit_size(paragraphs, w, h, start_pt=pt, space_after_pt=space_after,
                            indent_in=0.3 if (bullet or numbered) else 0.0)
        FIT_REPORT.append((title_for_report, name, pt, fits))
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = (f"{i + 1}.  " if numbered else "") + text
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
        if bullet:
            _set_bullet(p)
        elif numbered:                       # hanging indent for "1.  text"
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.45))))
            pPr.set("indent", str(-Emu(Inches(0.45))))
    return tb


def title_bar(slide, title, number, total):
    add_rect(slide, 0, 0, SLIDE_W, BAR_H, config.NAVY)
    add_rect(slide, 0, BAR_H, SLIDE_W, 0.08, config.ACCENT)
    pt = 36 if len(title) <= 40 else TITLE_PT
    FIT_REPORT.append((title, "title", pt, estimate_lines(title, SLIDE_W - 2.3, pt) == 1))
    add_text(slide, MARGIN, 0.08, SLIDE_W - 2.3, BAR_H - 0.16, [title], pt=pt, color="#FFFFFF", bold=True,
             anchor=MSO_ANCHOR.MIDDLE, autofit=False)
    add_text(slide, SLIDE_W - 1.8, 0.08, 1.4, BAR_H - 0.16, [f"{number} / {total}"], pt=18,
             color=config.ACCENT, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, autofit=False)


def source_tag(slide, src, x=MARGIN, y=BAR_H + 0.22):
    s = add_rect(slide, x, y, 3.2, 0.42, config.SOURCE_COLORS[src], MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = s.text_frame
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = src
    r.font.size, r.font.bold, r.font.name = Pt(16), True, FONT
    r.font.color.rgb = rgb("#FFFFFF")


def add_picture_fit(slide, path, x, y, max_w, max_h):
    """Insert a PNG scaled to fit inside the box, centred."""
    from PIL import Image
    with Image.open(path) as im:
        ar = im.width / im.height
    w = min(max_w, max_h * ar)
    h = w / ar
    slide.shapes.add_picture(str(path), Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2),
                             Inches(w), Inches(h))


def why_box(slide, points, x, y, w, h, title):
    add_rect(slide, x, y, w, h, "#FFF6E5")
    add_rect(slide, x, y, 0.12, h, config.ACCENT)
    add_text(slide, x + 0.25, y + 0.08, w - 0.35, 0.45, ["Why it matters for our dog"], pt=20,
             color=config.NAVY, bold=True, autofit=False)
    add_text(slide, x + 0.25, y + 0.55, w - 0.35, h - 0.62, points, pt=18, bullet=True,
             name="why", title_for_report=title)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------------
def run_demos():
    """Run every demo at default slider values, save PNGs, return metrics."""
    t0 = time.perf_counter()
    figs, metrics = {}, {}
    figs.update(pipeline.run()["figures"])
    act = actuator_net.run(config.ACT_DEFAULT_DELAY_MS)
    figs.update(act["figures"])
    metrics["actuator"] = {"delay_ms": act["delay_ms"], "rmse_Nm": act["errors"], "sweep": act["sweep"],
                           "seconds": round(act["seconds"], 2)}
    dr = domain_rand.run()
    figs.update(dr["figures"])
    metrics["domain_randomization"] = {"controllers_kp_kd_ff": dr["controllers"], "success_rates": dr["results"],
                                       "default_slider_outcome": dr["live_outcome"],
                                       "seconds": round(dr["seconds"], 2)}
    k2g = kine2go.run(download=True, with_gif=False)
    figs.update(k2g["figures"])
    metrics["kine2go"] = {"synthetic_fallback_used": k2g["synthetic"], "message": k2g["message"],
                          "shape": k2g["shape"], "seconds": round(k2g["seconds"], 2)}
    sdk = sdk_demo.run(with_gif=False)
    figs.update(sdk["figures"])
    metrics["sdk"] = {"move_commands": sdk["n_move_cmds"], "spots_investigated": sdk["n_spots"],
                      "policy_steps": len(sdk["loop"]["snapshots"]), "control_ticks": config.SDK_LOOP_TICKS,
                      "real_sdk_installed": sdk["real_sdk"], "seconds": round(sdk["seconds"], 2)}
    paths = {name: save_figure(fig, name) for name, fig in figs.items()}
    plt.close("all")
    metrics["total_demo_seconds"] = round(time.perf_counter() - t0, 2)
    return paths, metrics


def build_presentation(out_path=config.PPTX_PATH):
    FIT_REPORT.clear()
    paths, metrics = run_demos()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    S = {s["key"]: s for s in content.SLIDES}
    deck = []   # list of (builder function) so we know the total count first

    # ---- Title -------------------------------------------------------------
    def title_slide(sl, n, total):
        s = S["title"]
        add_rect(sl, 0, 0, SLIDE_W, SLIDE_H, config.NAVY)
        add_rect(sl, 0, 4.55, SLIDE_W, 0.1, config.ACCENT)
        add_text(sl, 0.8, 1.2, SLIDE_W - 1.6, 2.2, [s["title"]], pt=44, color="#FFFFFF", bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, autofit=False)
        add_text(sl, 0.8, 3.45, SLIDE_W - 1.6, 0.8, [s["subtitle"]], pt=24, color=config.ACCENT,
                 align=PP_ALIGN.CENTER, autofit=False)
        add_text(sl, 0.8, 4.95, SLIDE_W - 1.6, 1.9, [config.AUTHOR, config.COURSE, config.PRESENTATION_DATE],
                 pt=22, color="#FFFFFF", align=PP_ALIGN.CENTER, space_after=4, autofit=False)
        notes(sl, s["notes"])

    def pipeline_slide(sl, n, total):
        s = S["pipeline"]
        title_bar(sl, s["title"], n, total)
        add_picture_fit(sl, paths["pipeline"], MARGIN, BAR_H + 0.25, SLIDE_W - 2 * MARGIN, 4.35)
        add_text(sl, MARGIN + 0.3, 5.85, SLIDE_W - 2 * MARGIN - 0.3, 1.5, s["points"][:3], pt=18, bullet=True,
                 name="points", title_for_report=s["title"], space_after=4)
        notes(sl, s["notes"])

    def ideas_slide(key):
        def f(sl, n, total):
            s = S[key]
            title_bar(sl, s["title"], n, total)
            source_tag(sl, s["source"])
            add_text(sl, MARGIN, BAR_H + 0.75, SLIDE_W - 2 * MARGIN, SLIDE_H - BAR_H - 0.95, s["points"],
                     pt=19, bullet=True, name="key ideas", title_for_report=s["title"], space_after=7)
            notes(sl, s["notes"])
        return f

    def hwangbo_demo(sl, n, total):
        s = S["hwangbo"]
        title_bar(sl, "Demo: Toy Actuator Network", n, total)
        source_tag(sl, s["source"])
        add_picture_fit(sl, paths["actuator_comparison"], MARGIN, BAR_H + 0.7, 8.7, 4.3)
        e = metrics["actuator"]["rmse_Nm"]
        lines = [f"Hidden delay: {config.ACT_DEFAULT_DELAY_MS} ms (default)"] + \
                [f"{k.split(') ')[1]}: {v:.2f} N·m RMSE" for k, v in e.items()]
        add_text(sl, 9.45, BAR_H + 0.8, 3.45, 2.4, lines, pt=18, bullet=True, name="metrics",
                 title_for_report="hwangbo demo", space_after=6)
        add_text(sl, 9.45, BAR_H + 3.25, 3.45, 2.2, ["Toy model and synthetic logs. Not results from the paper."],
                 pt=16, color=config.ACCENT_3, bold=True, name="toy note", title_for_report="hwangbo demo")
        add_text(sl, MARGIN, 6.2, SLIDE_W - 2 * MARGIN, 1.2, [s["demo_caption"]], pt=16, color=config.GREY,
                 name="caption", title_for_report="hwangbo demo")
        notes(sl, s["notes"])

    def hwangbo_sweep(sl, n, total):
        s = S["hwangbo"]
        title_bar(sl, "Demo: Error Grows with Hidden Delay", n, total)
        source_tag(sl, s["source"])
        add_picture_fit(sl, paths["actuator_delay_sweep"], MARGIN, BAR_H + 0.7, 7.4, 3.2)
        sw = metrics["actuator"]["sweep"]
        lines = [f"Ideal model: {sw['ideal_rmse'][0]:.2f} N·m at 0 ms -> {sw['ideal_rmse'][-1]:.2f} N·m at 20 ms",
                 f"Analytical model: {sw['analytical_rmse'][0]:.2f} -> {sw['analytical_rmse'][-1]:.2f} N·m",
                 "The learned net sees a short history, so it can absorb the delay."]
        add_text(sl, 8.2, BAR_H + 0.8, 4.7, 3.9, lines, pt=18, bullet=True, name="sweep text",
                 title_for_report="hwangbo sweep")
        why_box(sl, s["why"], MARGIN, 5.15, SLIDE_W - 2 * MARGIN, 2.1, "hwangbo why")
        notes(sl, "Slider default is 10 ms. " + s["notes"])

    def survey_demo(sl, n, total):
        s = S["survey"]
        title_bar(sl, "Demo: Domain Randomization (Toy Cart)", n, total)
        source_tag(sl, s["source"])
        add_picture_fit(sl, paths["dr_success_rates"], MARGIN, BAR_H + 0.7, 8.6, 3.15)
        r = metrics["domain_randomization"]["success_rates"]
        names = list(r)
        lines = [f"Random mix: fixed {100 * r[names[0]]['Random mix']:.0f}% vs randomized "
                 f"{100 * r[names[1]]['Random mix']:.0f}%",
                 f"Carpet: fixed {100 * r[names[0]]['Carpet (high friction)']:.0f}% vs randomized "
                 f"{100 * r[names[1]]['Carpet (high friction)']:.0f}%",
                 "Simulated toy results only."]
        add_text(sl, 9.3, BAR_H + 0.8, 3.6, 3.6, lines, pt=18, bullet=True, name="dr text",
                 title_for_report="survey demo")
        why_box(sl, s["why"], MARGIN, 5.1, SLIDE_W - 2 * MARGIN, 2.2, "survey why")
        notes(sl, s["notes"])

    def kine2go_demo(sl, n, total):
        s = S["kine2go"]
        title_bar(sl, "Demo: Inside One Motion File", n, total)
        source_tag(sl, s["source"])
        syn = metrics["kine2go"]["synthetic_fallback_used"]
        label = ("SYNTHETIC fallback sample shown (download failed) - not Kine2Go data" if syn
                 else "Real Kine2Go motion.npy downloaded from Hugging Face")
        add_text(sl, 3.9, BAR_H + 0.2, 9.0, 0.5, [label], pt=16, bold=True,
                 color=config.ACCENT_3 if syn else "#1F9D55", autofit=False)
        cw, ch = 6.1, 2.75
        add_picture_fit(sl, paths["kine2go_joints"], MARGIN, BAR_H + 0.75, cw, ch)
        add_picture_fit(sl, paths["kine2go_feet"], MARGIN + cw + 0.2, BAR_H + 0.75, cw, ch)
        add_picture_fit(sl, paths["kine2go_footfall"], MARGIN, BAR_H + 0.8 + ch, cw, ch)
        why_box(sl, s["why"], MARGIN + cw + 0.3, BAR_H + 0.9 + ch, cw - 0.1, ch - 0.1, "kine2go why")
        notes(sl, s["notes"] + " " + s["demo_caption"])

    def sdk_demo_slide(sl, n, total):
        s = S["sdk"]
        title_bar(sl, "Demo: Mock SDK Policy Loop + Investigation", n, total)
        source_tag(sl, s["source"])
        add_text(sl, 3.9, BAR_H + 0.2, 9.0, 0.5, ["MOCK SDK - no hardware, no network; random untrained network"],
                 pt=16, bold=True, color=config.ACCENT_3, autofit=False)
        add_picture_fit(sl, paths["sdk_policy_loop"], MARGIN, BAR_H + 0.75, 7.6, 3.2)
        add_picture_fit(sl, paths["sdk_investigation"], 8.3, BAR_H + 0.7, 4.6, 3.3)
        why_box(sl, s["why"], MARGIN, 5.2, SLIDE_W - 2 * MARGIN, 2.1, "sdk why")
        notes(sl, s["notes"] + " " + s["demo_caption"])

    def summary_slide(sl, n, total):
        s = S["summary"]
        title_bar(sl, s["title"], n, total)
        rows, cols = len(s["table"]) + 1, 4
        widths = [2.9, 3.9, 2.6, 2.93]
        tbl = sl.shapes.add_table(rows, cols, Inches(MARGIN), Inches(BAR_H + 0.45), Inches(sum(widths)),
                                  Inches(5.2)).table
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(rows):
            for j in range(cols):
                cell = tbl.cell(i, j)
                text = s["table_header"][j] if i == 0 else s["table"][i - 1][j]
                cell.text = text
                p = cell.text_frame.paragraphs[0]
                p.runs[0].font.size = Pt(18 if i == 0 else 17)
                p.runs[0].font.bold = i == 0 or j == 0
                p.runs[0].font.name = FONT
                p.runs[0].font.color.rgb = rgb("#FFFFFF" if i == 0 else
                                               (config.SOURCE_COLORS[text] if j == 0 else config.TEXT_DARK))
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(config.NAVY if i == 0 else ("#FFFFFF" if i % 2 else config.BG_LIGHT))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                if i > 0:
                    FIT_REPORT.append(("summary table", f"cell {i},{j}", 17,
                                       estimate_lines(text, widths[j], 17) * 17 * 1.2 / 72 + 0.2 <= 5.2 / rows + 0.4))
        notes(sl, s["notes"])

    def next_steps_slide(sl, n, total):
        s = S["summary"]
        title_bar(sl, "Next Steps", n, total)
        add_text(sl, MARGIN + 0.2, BAR_H + 0.6, SLIDE_W - 2 * MARGIN - 0.4, 5.3, s["next_steps"], pt=24,
                 numbered=True, name="next steps", title_for_report="next steps", space_after=18)
        notes(sl, "Next steps follow the pipeline from left to right: simulation, training with randomization, "
                  "logging real Go2 data, and deployment through the SDK with safety limits.")

    def references_slide(sl, n, total):
        s = S["references"]
        title_bar(sl, s["title"], n, total)
        tb = add_text(sl, MARGIN + 0.2, BAR_H + 0.5, SLIDE_W - 2 * MARGIN - 0.4, 5.6, s["refs"], pt=20,
                      name="refs", title_for_report="references", space_after=14)
        for p in tb.text_frame.paragraphs:       # APA hanging indent
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.5))))
            pPr.set("indent", str(-Emu(Inches(0.5))))
        notes(sl, s["notes"])

    deck = [title_slide, pipeline_slide,
            ideas_slide("hwangbo"), hwangbo_demo, hwangbo_sweep,
            ideas_slide("survey"), survey_demo,
            ideas_slide("kine2go"), kine2go_demo,
            ideas_slide("sdk"), sdk_demo_slide,
            summary_slide, next_steps_slide, references_slide]
    for i, builder in enumerate(deck):
        builder(prs.slides.add_slide(blank), i + 1, len(deck))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    with open(config.OUTPUT_DIR / "demo_results.json", "w", encoding="utf-8") as fh:
        json.dump(metrics, fh, indent=2, default=lambda o: list(o) if isinstance(o, tuple) else str(o))
    return out_path


if __name__ == "__main__":
    t0 = time.perf_counter()
    path = build_presentation()
    print(f"Saved {path}  ({time.perf_counter() - t0:.1f} s)")
    bad = [r for r in FIT_REPORT if not r[3]]
    print(f"Text-fit check: {len(FIT_REPORT)} boxes, {len(bad)} possible overflows")
    for r in bad:
        print("  possible overflow:", r)
    print(f"Smallest body font used: {min(r[2] for r in FIT_REPORT)} pt")
